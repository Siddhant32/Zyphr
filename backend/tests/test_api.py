from pathlib import Path
from uuid import UUID

import httpx
import pytest
from docx import Document
from fastapi import HTTPException
from fastapi.testclient import TestClient

from app import main


@pytest.fixture
def client():
    with TestClient(main.app) as test_client:
        yield test_client
    main.app.dependency_overrides.clear()
    main.generation_requests.clear()


def test_admin_endpoints_require_authentication(client: TestClient) -> None:
    response = client.get("/api/admin/stats")
    assert response.status_code == 401


def test_public_library_queries_published_documents_only(
    client: TestClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    async def fake_request(method, path, **kwargs):
        assert method == "GET"
        assert path == "/rest/v1/documents"
        assert kwargs["params"]["published"] == "eq.true"
        return httpx.Response(
            200,
            json=[
                {
                    "id": "a",
                    "title": "DBMS",
                    "subject": "Database",
                    "description": "",
                    "cover_image_path": None,
                    "created_at": "2026-01-01T00:00:00Z",
                }
            ],
        )

    monkeypatch.setattr(main, "_supabase_request", fake_request)
    response = client.get("/api/practicals")

    assert response.status_code == 200
    assert response.json()["items"][0]["title"] == "DBMS"
    assert "original_file_path" not in response.json()["items"][0]


def test_public_practical_includes_configured_student_fields(
    client: TestClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    async def fake_load(identifier, *, published_only=False):
        assert str(identifier) == "7e3be4a8-211b-4f55-bd2c-9e19723973d3"
        assert published_only
        return {
            "id": str(identifier),
            "title": "Certificate",
            "extra_fields": ["branch", "year"],
            "published": True,
        }

    monkeypatch.setattr(main, "_load_document", fake_load)
    response = client.get("/api/practicals/7e3be4a8-211b-4f55-bd2c-9e19723973d3")

    assert response.status_code == 200
    assert response.json()["extra_fields"] == ["branch", "year"]


def test_generation_uses_fresh_copy_and_does_not_change_source(
    client: TestClient, monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    document_id = UUID("7e3be4a8-211b-4f55-bd2c-9e19723973d3")
    original_path = tmp_path / "original.docx"
    source_document = Document()
    source_document.add_paragraph("Name of Student: Old Student")
    source_document.add_paragraph("Reg ID: OLD-001")
    source_document.add_paragraph("Branch: Old Branch")
    source_document.save(original_path)
    source_bytes = original_path.read_bytes()
    personalized_text: list[str] = []

    async def fake_load(identifier, *, published_only=False):
        assert identifier == document_id
        assert published_only
        return {
            "id": str(document_id),
            "published": True,
            "original_file_path": f"documents/{document_id}.docx",
            "extra_fields": ["branch"],
        }

    async def fake_request(method, path, **kwargs):
        if path.startswith("/storage/v1/object/"):
            return httpx.Response(200, content=source_bytes)
        if path.endswith("/increment_generation_count"):
            return httpx.Response(204)
        raise AssertionError(f"Unexpected Supabase request: {method} {path}")

    def fake_convert_docx(source: Path, output_directory: Path, timeout: int) -> Path:
        personalized_text.append("\n".join(p.text for p in Document(source).paragraphs))
        pdf = output_directory / f"{source.stem}.pdf"
        pdf.write_bytes(b"%PDF-1.7\nmock pdf\n")
        return pdf

    monkeypatch.setattr(main, "_load_document", fake_load)
    monkeypatch.setattr(main, "_supabase_request", fake_request)
    monkeypatch.setattr(main, "convert_docx_to_pdf", fake_convert_docx)

    first = client.post(
        "/api/generate",
        json={
            "document_id": str(document_id),
            "full_name": "Asha Rao",
            "registration_id": "00-CS-81",
            "extra_fields": {"branch": "Computer Science"},
        },
    )
    second = client.post(
        "/api/generate",
        json={
            "document_id": str(document_id),
            "full_name": "Mira Sen",
            "registration_id": "00-CS-82",
            "extra_fields": {"branch": "Information Technology"},
        },
    )

    assert first.status_code == second.status_code == 200
    assert first.headers["cache-control"] == "no-store"
    assert first.content.startswith(b"%PDF-")
    assert "Asha Rao" in personalized_text[0]
    assert "00-CS-81" in personalized_text[0]
    assert "Computer Science" in personalized_text[0]
    assert "Mira Sen" in personalized_text[1]
    assert "00-CS-82" in personalized_text[1]
    assert "Information Technology" in personalized_text[1]
    assert "Asha Rao" not in personalized_text[1]
    assert original_path.read_bytes() == source_bytes


def test_generation_rejects_unpublished_documents(
    client: TestClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    async def unpublished(identifier, *, published_only=False):
        raise HTTPException(status_code=404, detail="Practical not found.")

    monkeypatch.setattr(main, "_load_document", unpublished)
    response = client.post(
        "/api/generate",
        json={
            "document_id": "7e3be4a8-211b-4f55-bd2c-9e19723973d3",
            "full_name": "Asha Rao",
            "registration_id": "0012",
        },
    )
    assert response.status_code == 404


def test_generation_requires_all_configured_extra_fields(
    client: TestClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    async def configured_document(identifier, *, published_only=False):
        return {"id": str(identifier), "published": True, "extra_fields": ["branch"]}

    monkeypatch.setattr(main, "_load_document", configured_document)
    response = client.post(
        "/api/generate",
        json={
            "document_id": "7e3be4a8-211b-4f55-bd2c-9e19723973d3",
            "full_name": "Asha Rao",
            "registration_id": "0012",
        },
    )
    assert response.status_code == 422
    assert "every additional field" in response.json()["detail"]


def test_admin_stats_and_list_hide_source_storage_paths(
    client: TestClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    main.app.dependency_overrides[main._require_admin] = lambda: "admin-user"

    async def fake_request(method, path, **kwargs):
        if path.endswith("/get_zyphr_dashboard_stats"):
            return httpx.Response(
                200,
                json={
                    "total_documents": 1,
                    "published_documents": 1,
                    "unpublished_documents": 0,
                    "total_generations": 3,
                },
            )
        return httpx.Response(
            200,
            json=[
                {
                    "id": "doc-id",
                    "title": "Operating Systems",
                    "subject": "OS",
                    "description": "",
                    "original_file_path": "documents/private.docx",
                    "cover_image_path": None,
                    "published": True,
                    "generation_count": 3,
                }
            ],
        )

    monkeypatch.setattr(main, "_supabase_request", fake_request)
    stats = client.get("/api/admin/stats")
    documents = client.get("/api/admin/documents")

    assert stats.json()["total_generations"] == 3
    assert documents.status_code == 200
    assert "original_file_path" not in documents.json()["items"][0]


def test_admin_upload_validates_and_stores_docx(
    client: TestClient, monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    main.app.dependency_overrides[main._require_admin] = lambda: "admin-user"
    source = tmp_path / "practical.docx"
    document = Document()
    document.add_paragraph("Name of Student:")
    document.save(source)
    uploads: list[str] = []

    async def fake_upload(bucket, path, data, content_type):
        uploads.append(path)

    async def fake_request(method, path, **kwargs):
        body = kwargs["json"]
        return httpx.Response(
            201,
            json=[
                {
                    **body,
                    "published": False,
                    "generation_count": 0,
                    "created_at": "2026-01-01T00:00:00Z",
                    "updated_at": "2026-01-01T00:00:00Z",
                }
            ],
        )

    monkeypatch.setattr(main, "_upload_storage", fake_upload)
    monkeypatch.setattr(main, "_supabase_request", fake_request)
    monkeypatch.setattr(main, "_service_headers", lambda extra=None: extra or {})
    with source.open("rb") as file:
        response = client.post(
            "/api/admin/documents",
            files={
                "file": (
                    "practical.docx",
                    file,
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )
            },
            data={
                "title": "Operating Systems",
                "subject": "OS",
                "description": "Practical exercises",
                "extra_fields": ["branch"],
            },
        )

    assert response.status_code == 201
    assert len(uploads) == 1
    assert response.json()["title"] == "Operating Systems"
    assert response.json()["extra_fields"] == ["branch"]
    assert "original_file_path" not in response.json()


def test_admin_upload_rejects_corrupt_docx(
    client: TestClient,
) -> None:
    main.app.dependency_overrides[main._require_admin] = lambda: "admin-user"
    response = client.post(
        "/api/admin/documents",
        files={"file": ("broken.docx", b"not a Word file", "application/octet-stream")},
        data={"title": "Broken", "subject": "OS", "description": "Bad file"},
    )
    assert response.status_code == 422
    assert "corrupt" in response.json()["detail"]
