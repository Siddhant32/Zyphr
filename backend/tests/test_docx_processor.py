from pathlib import Path

import pytest
from docx import Document

from app.docx_processor import (
    _libreoffice_executable,
    convert_docx_to_pdf,
    personalize_docx,
)


def test_replaces_split_runs_and_multiple_experiment_pages(tmp_path: Path) -> None:
    source = tmp_path / "source.docx"
    output = tmp_path / "personalized.docx"
    document = Document()
    paragraph = document.add_paragraph()
    paragraph.add_run("Name of ")
    paragraph.add_run("Student: ")
    paragraph.add_run("Previous Name").bold = True
    document.add_page_break()
    document.add_paragraph("Reg. Id.: 0000-A")
    document.add_page_break()
    document.add_paragraph("Name of Students")
    document.save(source)

    counts = personalize_docx(source, output, "Asha Rao", "REG-0081")
    result = Document(output)
    paragraphs = [paragraph.text for paragraph in result.paragraphs]

    assert counts == {"name": 2, "registration_id": 1}
    assert paragraphs[0] == "Name of Student: Asha Rao"
    assert paragraphs[2] == "Reg. Id.: REG-0081"
    assert paragraphs[4] == "Name of Students Asha Rao"
    assert Document(source).paragraphs[0].text == "Name of Student: Previous Name"


def test_fills_table_neighbor_and_header_footer_fields(tmp_path: Path) -> None:
    source = tmp_path / "table.docx"
    output = tmp_path / "table-filled.docx"
    document = Document()
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Student Name:"
    table.cell(1, 0).text = "Registration No.:"
    document.sections[0].header.paragraphs[0].text = "Name: old value"
    document.sections[0].footer.paragraphs[0].text = "Roll No."
    document.save(source)

    counts = personalize_docx(source, output, "Mira Sen", "00-12A")
    result = Document(output)

    assert counts == {"name": 2, "registration_id": 2}
    assert result.tables[0].cell(0, 1).text == "Mira Sen"
    assert result.tables[0].cell(1, 1).text == "00-12A"
    assert result.sections[0].header.paragraphs[0].text == "Name: Mira Sen"
    assert result.sections[0].footer.paragraphs[0].text == "Roll No.00-12A"


def test_replaces_selected_additional_student_fields(tmp_path: Path) -> None:
    source = tmp_path / "additional-fields.docx"
    output = tmp_path / "additional-fields-filled.docx"
    document = Document()
    document.add_paragraph("Name: ")
    document.add_paragraph("Registration ID: ")
    document.add_paragraph("Branch: Old Branch")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Department:"
    document.save(source)

    counts = personalize_docx(
        source,
        output,
        "Asha Rao",
        "REG-0081",
        {"branch": "Computer Science", "department": "Engineering"},
    )
    result = Document(output)

    assert counts == {
        "name": 1,
        "registration_id": 1,
        "branch": 1,
        "department": 1,
    }
    assert [paragraph.text for paragraph in result.paragraphs] == [
        "Name: Asha Rao",
        "Registration ID: REG-0081",
        "Branch: Computer Science",
    ]
    assert result.tables[0].cell(0, 1).text == "Engineering"


def test_requires_a_docx_label_for_each_selected_extra_field(tmp_path: Path) -> None:
    source = tmp_path / "missing-branch.docx"
    output = tmp_path / "should-not-exist.docx"
    document = Document()
    document.add_paragraph("Name: ")
    document.add_paragraph("Registration ID: ")
    document.save(source)

    with pytest.raises(ValueError, match="branch"):
        personalize_docx(source, output, "Asha Rao", "REG-0081", {"branch": "CSE"})
    assert not output.exists()


def test_personalizes_certificate_template_placeholders_without_losing_text(
    tmp_path: Path,
) -> None:
    source = tmp_path / "certificate.docx"
    output = tmp_path / "certificate-filled.docx"
    document = Document()
    document.add_paragraph(
        "This is to certify that Mr./Ms. "
        "____________________________________________________________"
    )
    document.add_paragraph(
        "(Reg. ID ________________________) from   B. Tech "
        "__________________________________________ has"
    )
    document.add_paragraph(
        "Academic Year _____2026-2027______. The student has satisfactorily "
        "performed all the required experiments."
    )
    document.save(source)

    counts = personalize_docx(
        source,
        output,
        "Siddhant Chandanshiv",
        "UST26L0102051",
        {"branch": "Computer Science", "year": "2026-2027"},
    )
    result = Document(output)
    paragraphs = [paragraph.text for paragraph in result.paragraphs]

    assert counts == {"name": 1, "registration_id": 1, "branch": 1, "year": 1}
    assert "Mr./Ms. Siddhant Chandanshiv" in paragraphs[0]
    assert "(Reg. ID UST26L0102051) from   B. Tech Computer Science has" in paragraphs[1]
    assert "Academic Year 2026-2027. The student has satisfactorily" in paragraphs[2]


def test_rejects_documents_without_recognized_student_fields(tmp_path: Path) -> None:
    source = tmp_path / "no-fields.docx"
    output = tmp_path / "should-not-exist.docx"
    document = Document()
    document.add_paragraph("Operating Systems Practical")
    document.save(source)

    with pytest.raises(ValueError, match="No recognizable"):
        personalize_docx(source, output, "Asha Rao", "REG-0081")
    assert not output.exists()


def test_converts_docx_to_pdf_when_libreoffice_is_installed(tmp_path: Path) -> None:
    try:
        _libreoffice_executable()
    except RuntimeError:
        pytest.skip("LibreOffice is not installed.")

    source = tmp_path / "source.docx"
    document = Document()
    document.add_paragraph("PDF conversion smoke test")
    document.save(source)

    pdf = convert_docx_to_pdf(source, tmp_path, timeout_seconds=45)

    assert pdf.read_bytes().startswith(b"%PDF-")
