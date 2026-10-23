import os
import re
import shutil
import subprocess
import uuid
from pathlib import Path
from typing import Iterable

from docx import Document
from docx.document import Document as DocumentObject
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P


FIELD_LABELS: dict[str, tuple[str, ...]] = {
    "name": (
        r"name\s+of\s+students?",
        r"student\s+name",
        r"name",
    ),
    "registration_id": (
        r"registration\s+(?:id|no\.?)",
        r"reg\.?\s*(?:id|no\.?)",
        r"roll\s+no\.?",
    ),
    "branch": (r"branch(?:\s+name)?", r"stream"),
    "department": (r"department", r"dept\.?"),
    "year": (r"academic\s+year", r"year\s+of\s+study", r"year"),
    "section": (r"class\s+section", r"section"),
}
EXTRA_FIELD_IDS = frozenset({"branch", "department", "year", "section"})
FIELD_DISPLAY_NAMES = {
    "name": "student name",
    "registration_id": "registration ID",
    "branch": "branch",
    "department": "department",
    "year": "year",
    "section": "section",
}

LABEL_PATTERN = re.compile(
    r"^\s*(?P<label>"
    + "|".join(
        f"(?:{label})" for field_labels in FIELD_LABELS.values() for label in field_labels
    )
    + r")\s*(?P<separator>[:：.\-]*\s*)",
    re.IGNORECASE,
)


def _libreoffice_executable() -> str:
    for name in ("libreoffice", "soffice"):
        executable = shutil.which(name)
        if executable:
            return executable

    if os.name == "nt":
        for program_files in (
            os.environ.get("ProgramFiles"),
            os.environ.get("ProgramFiles(x86)"),
        ):
            if program_files:
                executable = (
                    Path(program_files) / "LibreOffice" / "program" / "soffice.exe"
                )
                if executable.is_file():
                    return str(executable)

    raise RuntimeError("LibreOffice is not installed or could not be found.")


def _paragraphs_in_container(container: DocumentObject | _Cell) -> Iterable[Paragraph]:
    element_container = (
        container._element.body
        if isinstance(container, DocumentObject)
        else container._element
    )
    for element in element_container.iterchildren():
        if isinstance(element, CT_P):
            yield Paragraph(element, container)
        elif isinstance(element, CT_Tbl):
            table = Table(element, container)
            for row in table.rows:
                for cell in row.cells:
                    yield from _paragraphs_in_container(cell)


def _all_paragraphs(document: DocumentObject) -> Iterable[Paragraph]:
    yield from _paragraphs_in_container(document)
    for section in document.sections:
        for part in (
            section.header,
            section.footer,
            section.first_page_header,
            section.first_page_footer,
            section.even_page_header,
            section.even_page_footer,
        ):
            yield from _paragraphs_in_container(part)


def _replace_span(paragraph: Paragraph, start: int, end: int, value: str) -> None:
    runs = paragraph.runs
    text = "".join(run.text for run in runs)
    if start > len(text) or end > len(text):
        raise ValueError("Student field could not be safely mapped to document text.")

    offsets: list[tuple[int, int]] = []
    cursor = 0
    for run in runs:
        offsets.append((cursor, cursor + len(run.text)))
        cursor += len(run.text)

    touched = [
        index
        for index, (left, right) in enumerate(offsets)
        if right > start and left < end
    ]
    if not touched:
        if start == end and runs:
            runs[-1].text += value
            return
        if not runs:
            paragraph.add_run(value)
            return
        raise ValueError("Student field could not be safely mapped to document runs.")

    first, last = touched[0], touched[-1]
    first_left, first_right = offsets[first]
    last_left, last_right = offsets[last]
    if first == last:
        run = runs[first]
        run.text = (
            run.text[: start - first_left]
            + value
            + run.text[end - first_left :]
        )
        return

    runs[first].text = runs[first].text[: start - first_left] + value
    for index in touched[1:-1]:
        runs[index].text = ""
    runs[last].text = runs[last].text[end - last_left :]
    if last_right < end:
        raise ValueError("Student field could not be safely mapped to document runs.")


def _replace_field_in_paragraph(paragraph: Paragraph, field: str, value: str) -> int:
    text = paragraph.text
    match = LABEL_PATTERN.match(text)
    if not match:
        return 0

    label = match.group("label").casefold()
    field_for_label = next(
        (
            field_name
            for field_name, patterns in FIELD_LABELS.items()
            if any(re.fullmatch(pattern, label, flags=re.IGNORECASE) for pattern in patterns)
        ),
        None,
    )
    if field_for_label != field:
        return 0

    value_start = match.end()
    value_end = len(text)
    if text[value_start:value_end].strip() == value:
        return 0
    replacement = value
    if (
        value_start == value_end
        and not text.endswith((" ", "\t", ":", "：", ".", "-"))
    ):
        replacement = f" {value}"
    _replace_span(paragraph, value_start, value_end, replacement)
    return 1


def _set_empty_cell(cell: _Cell, value: str) -> None:
    if len(cell.paragraphs) == 1 and not cell.paragraphs[0].text:
        cell.paragraphs[0].add_run(value)
    else:
        cell.add_paragraph(value)


def _replace_table_neighbor_fields(
    document: DocumentObject, values: dict[str, str]
) -> tuple[dict[str, int], set[object]]:
    replacements = {field: 0 for field in values}
    handled_labels: set[object] = set()
    for table in document.tables:
        for row in table.rows:
            for cell_index, cell in enumerate(row.cells[:-1]):
                label_match = LABEL_PATTERN.match(cell.text)
                if not label_match or cell_index + 1 >= len(row.cells):
                    continue
                label = label_match.group("label").casefold()
                field = next(
                    (
                        field_name
                        for field_name, patterns in FIELD_LABELS.items()
                        if any(
                            re.fullmatch(pattern, label, flags=re.IGNORECASE)
                            for pattern in patterns
                        )
                    ),
                    None,
                )
                target = row.cells[cell_index + 1]
                if field and not target.text.strip():
                    _set_empty_cell(target, values[field])
                    replacements[field] += 1
                    handled_labels.update(
                        paragraph._p
                        for paragraph in cell.paragraphs
                        if LABEL_PATTERN.match(paragraph.text)
                    )
    return replacements, handled_labels


def personalize_docx(
    source: Path,
    destination: Path,
    name: str,
    registration_id: str,
    extra_fields: dict[str, str] | None = None,
) -> dict[str, int]:
    document = Document(source)
    values = {"name": name, "registration_id": registration_id, **(extra_fields or {})}
    unsupported_fields = values.keys() - FIELD_LABELS.keys()
    if unsupported_fields:
        raise ValueError("One or more selected student detail fields are unsupported.")
    counts = dict.fromkeys(values, 0)
    neighboring_counts, handled_labels = _replace_table_neighbor_fields(
        document, values
    )

    for paragraph in _all_paragraphs(document):
        if paragraph._p in handled_labels:
            continue
        for field in values:
            counts[field] += _replace_field_in_paragraph(
                paragraph, field, values[field]
            )

    for field in values:
        counts[field] += neighboring_counts[field]
    missing_fields = [
        FIELD_DISPLAY_NAMES[field]
        for field, count in counts.items()
        if field in (extra_fields or {}) and not count
    ]
    if missing_fields:
        missing = ", ".join(missing_fields)
        raise ValueError(
            f"No recognizable DOCX field was found for {missing}. "
            "Ask the administrator to add a supported label to the DOCX."
        )
    if not sum(counts.values()):
        raise ValueError(
            "No recognizable student name or registration ID fields were found. "
            "Ask the administrator to add a supported label to the DOCX."
        )

    destination.parent.mkdir(parents=True, exist_ok=True)
    document.save(destination)
    Document(destination)
    return counts


def convert_docx_to_pdf(source: Path, output_directory: Path, timeout_seconds: int) -> Path:
    output_directory.mkdir(parents=True, exist_ok=True)
    profile = output_directory / f"lo-profile-{uuid.uuid4().hex}"
    result = subprocess.run(
        [
            _libreoffice_executable(),
            f"-env:UserInstallation={profile.resolve().as_uri()}",
            "--headless",
            "--convert-to",
            "pdf",
            "--outdir",
            str(output_directory),
            str(source),
        ],
        capture_output=True,
        text=True,
        timeout=timeout_seconds,
        check=False,
    )
    pdf_path = output_directory / f"{source.stem}.pdf"
    if result.returncode != 0 or not pdf_path.is_file() or pdf_path.stat().st_size < 5:
        raise RuntimeError("The document could not be converted to PDF.")
    if pdf_path.read_bytes()[:5] != b"%PDF-":
        raise RuntimeError("The document conversion did not produce a valid PDF.")
    return pdf_path
